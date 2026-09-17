"""
routers/asistencias.py — Endpoint del Lector y Gestión de Asistencias

Este archivo contiene:
  1. POST /api/scan : Recibe los escaneos del lector físico (profesor o alumno).
  2. GET /api/asistencias/{clase_id}/excel : Descarga el reporte de asistencia en Excel.
"""

from datetime import date, datetime
import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from sqlmodel import Session, select

from app.database import get_session
from app.models import Alumno, Asistencia, Clase, Profesor, Sesion

# Creamos el Router. main.py lo conectará con prefijo "/api"
router = APIRouter()


# ─────────────────────────────────────────────────────────────
# 1. Esquema de datos para la petición del escáner
# ─────────────────────────────────────────────────────────────
class ScanRequest(BaseModel):
    """
    Define la estructura que el lector (o script lector.py) debe enviar por HTTP.
    Ejemplo de JSON: { "codigo": "218123456" }
    """
    codigo: str


# ─────────────────────────────────────────────────────────────
# 2. Estado en memoria: Profesor en espera
# ─────────────────────────────────────────────────────────────
# Variable global temporal para recordar si un profesor acaba de escanear.
# Guarda un diccionario: {"profesor_id": 1, "nombre": "Juan Perez", "fecha": date(2026, 9, 1)}
profesor_en_espera: Optional[dict] = None


# ─────────────────────────────────────────────────────────────
# 3. Endpoint principal: POST /api/scan
# ─────────────────────────────────────────────────────────────
@router.post("/scan")
def procesar_escaneo(datos: ScanRequest, session: Session = Depends(get_session)):
    """
    Procesa el código leído por el lector de código de barras.
    Identifica automáticamente si es un Profesor o un Alumno y ejecuta la lógica.
    """
    global profesor_en_espera
    codigo_limpio = datos.codigo.strip()
    hoy = date.today()
    hora_actual = datetime.now().time()

    # ─────────────────────────────────────────────────────────
    # PASO A: ¿El código escaneado es de un PROFESOR?
    # Equivalente a: SELECT * FROM profesor WHERE codigo_profesor = '...'
    # ─────────────────────────────────────────────────────────
    stmt_profesor = select(Profesor).where(Profesor.codigo_profesor == codigo_limpio)
    profesor = session.exec(stmt_profesor).first()

    if profesor:
        # Ponemos al profesor en modo espera
        profesor_en_espera = {
            "profesor_id": profesor.id,
            "nombre": profesor.nombre_profesor,
            "fecha": hoy,
        }
        return {
            "status": "ok",
            "tipo": "profesor",
            "nombre": profesor.nombre_profesor,
            "mensaje": f"Profesor {profesor.nombre_profesor} listo. Esperando escaneo del primer alumno para iniciar clase.",
        }

    # ─────────────────────────────────────────────────────────
    # PASO B: ¿El código escaneado es de un ALUMNO?
    # Un alumno puede estar inscrito en más de una clase (materias diferentes).
    # ─────────────────────────────────────────────────────────
    stmt_alumnos = select(Alumno).where(Alumno.codigo_alumno == codigo_limpio)
    alumnos_encontrados = session.exec(stmt_alumnos).all()

    if alumnos_encontrados:
        alumno_elegido: Optional[Alumno] = None
        sesion_hoy: Optional[Sesion] = None
        clase: Optional[Clase] = None

        # 1. Si hay un profesor en espera, buscamos la clase de este profesor en la que esté el alumno
        if profesor_en_espera and profesor_en_espera["fecha"] == hoy:
            for al in alumnos_encontrados:
                c = session.get(Clase, al.clase_id)
                if c and c.profesor_id == profesor_en_espera["profesor_id"]:
                    alumno_elegido = al
                    clase = c
                    break

            if not alumno_elegido:
                return {
                    "status": "error",
                    "tipo": "alumno",
                    "mensaje": f"El alumno no está registrado en ninguna clase del profesor en espera ({profesor_en_espera['nombre']}).",
                }

            # Abrimos la sesión del día para esta clase
            stmt_sesion = select(Sesion).where(
                Sesion.clase_id == alumno_elegido.clase_id,
                Sesion.fecha == hoy,
            )
            sesion_hoy = session.exec(stmt_sesion).first()
            if not sesion_hoy:
                sesion_hoy = Sesion(clase_id=alumno_elegido.clase_id, fecha=hoy)
                session.add(sesion_hoy)
                session.commit()
                session.refresh(sesion_hoy)

            # Limpiamos el modo espera del profesor porque la clase ya inició
            profesor_en_espera = None

        else:
            # 2. Si no hay profesor en espera, buscamos si hay una sesión ya iniciada hoy para alguna clase del alumno
            for al in alumnos_encontrados:
                stmt_sesion = select(Sesion).where(
                    Sesion.clase_id == al.clase_id,
                    Sesion.fecha == hoy,
                )
                s = session.exec(stmt_sesion).first()
                if s:
                    alumno_elegido = al
                    sesion_hoy = s
                    clase = session.get(Clase, al.clase_id)
                    break

            if not sesion_hoy or not alumno_elegido:
                return {
                    "status": "error",
                    "tipo": "alumno",
                    "mensaje": "No hay sesión abierta hoy para la clase de este alumno. El profesor debe escanear su credencial primero.",
                }

        # 3. Llegados a este punto, tenemos un alumno y una sesión válida (sesion_hoy).
        # Verificamos si el alumno ya había registrado asistencia hoy
        stmt_asistencia = select(Asistencia).where(
            Asistencia.sesion_id == sesion_hoy.id,
            Asistencia.alumno_id == alumno_elegido.id,
        )
        asistencia_existente = session.exec(stmt_asistencia).first()

        if asistencia_existente:
            # Si ya escaneó antes, solo actualizamos la hora (no duplicamos registros)
            asistencia_existente.hora_llegada = hora_actual
            session.add(asistencia_existente)
            session.commit()
            return {
                "status": "ok",
                "tipo": "alumno",
                "nombre": alumno_elegido.nombre_alumno,
                "clase": clase.nombre_clase if clase else "",
                "mensaje": f"Asistencia ya registrada previamente para {alumno_elegido.nombre_alumno}. Hora actualizada.",
            }
        else:
            # Nuevo registro de asistencia
            nueva_asistencia = Asistencia(
                sesion_id=sesion_hoy.id,
                alumno_id=alumno_elegido.id,
                hora_llegada=hora_actual,
            )
            session.add(nueva_asistencia)
            session.commit()
            return {
                "status": "ok",
                "tipo": "alumno",
                "nombre": alumno_elegido.nombre_alumno,
                "clase": clase.nombre_clase if clase else "",
                "mensaje": f"Asistencia registrada con éxito: {alumno_elegido.nombre_alumno}",
            }

    # ─────────────────────────────────────────────────────────
    # PASO C: El código no es ni de Profesor ni de Alumno
    # ─────────────────────────────────────────────────────────
    return {
        "status": "error",
        "tipo": "desconocido",
        "mensaje": f"El código '{codigo_limpio}' no está registrado en el sistema.",
    }


# ─────────────────────────────────────────────────────────────
# 4. Endpoint: Descarga de Excel de Asistencias
# ─────────────────────────────────────────────────────────────
@router.get("/asistencias/{clase_id}/excel")
def descargar_excel_asistencias(clase_id: int, session: Session = Depends(get_session)):
    """
    Genera y descarga un archivo Excel (.xlsx) con la matriz de asistencia:
      - Filas: Alumnos de la clase (código y nombre)
      - Columnas: Fechas de las sesiones impartidas
      - Celdas: '✓' si asistió o '—' si faltó
    """
    clase = session.get(Clase, clase_id)
    if not clase:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Clase no encontrada")

    # Obtenemos los alumnos de la clase ordenados alfabéticamente
    stmt_alumnos = select(Alumno).where(Alumno.clase_id == clase_id).order_by(Alumno.nombre_alumno)
    alumnos = session.exec(stmt_alumnos).all()

    # Obtenemos todas las sesiones registradas para esta clase ordenadas por fecha
    stmt_sesiones = select(Sesion).where(Sesion.clase_id == clase_id).order_by(Sesion.fecha)
    sesiones = session.exec(stmt_sesiones).all()

    # Obtenemos todos los registros de asistencias para las sesiones de esta clase
    ids_sesiones = [s.id for s in sesiones if s.id is not None]
    asistencias_map = set()  # Guardará tuplas (sesion_id, alumno_id) para búsqueda rápida

    if ids_sesiones:
        stmt_asistencias = select(Asistencia).where(Asistencia.sesion_id.in_(ids_sesiones))
        for a in session.exec(stmt_asistencias).all():
            asistencias_map.add((a.sesion_id, a.alumno_id))

    # Creamos el libro de Excel con openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista de Asistencia"

    # Estilos visuales para el Excel
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center_align = Alignment(horizontal="center", vertical="center")

    # Encabezados: Código, Nombre de Alumno, y luego cada Fecha
    headers = ["Código", "Nombre del Alumno"]
    for sesion in sesiones:
        headers.append(sesion.fecha.strftime("%d/%m/%Y"))
    headers.append("Total Asistencias")

    ws.append(headers)

    # Aplicamos estilo a la fila de encabezados
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Llenamos los datos de cada alumno
    for alumno in alumnos:
        fila = [alumno.codigo_alumno, alumno.nombre_alumno]
        total_alumno = 0

        for sesion in sesiones:
            if (sesion.id, alumno.id) in asistencias_map:
                fila.append("✓")
                total_alumno += 1
            else:
                fila.append("—")

        fila.append(total_alumno)
        ws.append(fila)

    # Ajustamos el ancho de las columnas automáticamente
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Guardamos el archivo en memoria (sin guardarlo en disco físico)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    nombre_archivo = f"Asistencia_{clase.nombre_clase.replace(' ', '_')}_{date.today()}.xlsx"

    # Enviamos el archivo como descarga al navegador
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )